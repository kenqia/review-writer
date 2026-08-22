from __future__ import annotations

import base64
import hashlib
import json
import threading
import tempfile
from pathlib import Path
from typing import Any
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from jsonschema import Draft202012Validator
from docx import Document

from review_writer.project.manuscript_v2 import (
    approve_section,
    merge_authoritative_manuscript,
    register_section_draft,
)
from review_writer.project.paper_evidence import (
    apply_paper_evidence_decision,
    paper_evidence_state,
    register_paper_evidence_candidates,
)
from review_writer.project.parse_quality import (
    apply_parse_quality_decision,
    project_parse_quality_state,
    write_parse_quality_gate,
)
from review_writer.project.review_figures import (
    _content_list_v2_digest,
    _source_truth_digest,
)
from review_writer.project.section_contract import (
    apply_section_contract_decision,
    register_section_contracts,
    section_contract_state,
)
from review_writer.project.source_truth import (
    canonical_digest,
    load_source_truth_bundle,
    write_source_truth_bundle,
)
from review_writer.project.synthesis import (
    apply_comparison_protocol_decision,
    apply_synthesis_decision,
    register_comparison_protocol,
    register_synthesis_candidates,
    synthesis_state,
)
from view import serve_review_dashboard as dashboard


PROJECT_ID = "prod006-tiny-source-to-release"
STUDY_ID = "study-prod006"
SOURCE_ID = "source-prod006-main"
EVIDENCE_ID = "evidence-prod006-observation"
SYNTHESIS_ID = "synthesis-prod006-bounded"
SECTION_ID = "reported-result"

_TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
    "+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def _http_json(base_url: str, path: str) -> tuple[int, Any]:
    request = Request(f"{base_url}{path}", method="GET")
    try:
        with urlopen(request, timeout=15) as response:
            return response.status, json.loads(response.read().decode("utf-8"))
    except HTTPError as exc:  # pragma: no cover - keeps the failure body visible
        body = exc.read().decode("utf-8", errors="replace")
        raise AssertionError(f"GET {path} returned HTTP {exc.code}: {body}") from exc


def _http_text(base_url: str, path: str) -> tuple[int, str]:
    request = Request(f"{base_url}{path}", method="GET")
    try:
        with urlopen(request, timeout=15) as response:
            return response.status, response.read().decode("utf-8")
    except HTTPError as exc:  # pragma: no cover - keeps the failure body visible
        body = exc.read().decode("utf-8", errors="replace")
        raise AssertionError(f"GET {path} returned HTTP {exc.code}: {body}") from exc


def _http_bytes(base_url: str, path: str) -> tuple[int, bytes]:
    request = Request(f"{base_url}{path}", method="GET")
    try:
        with urlopen(request, timeout=15) as response:
            return response.status, response.read()
    except HTTPError as exc:  # pragma: no cover - keeps the failure body visible
        body = exc.read().decode("utf-8", errors="replace")
        raise AssertionError(f"GET {path} returned HTTP {exc.code}: {body}") from exc


def _post_release(base_url: str) -> tuple[int, dict[str, Any]]:
    # This byte string is intentionally exact: the Dashboard contract rejects
    # extra keys and any release level other than SELF_REVIEWED_DRAFT.
    body = b'{"release_level":"SELF_REVIEWED_DRAFT"}'
    request = Request(
        f"{base_url}/api/project/{PROJECT_ID}/export-docx",
        data=body,
        headers={
            "Content-Type": "application/json",
            "Content-Length": str(len(body)),
        },
        method="POST",
    )
    try:
        with urlopen(request, timeout=30) as response:
            return response.status, json.loads(response.read().decode("utf-8"))
    except HTTPError as exc:
        body_text = exc.read().decode("utf-8", errors="replace")
        raise AssertionError(f"POST export-docx returned HTTP {exc.code}: {body_text}") from exc


def _start_dashboard(review_root: Path) -> tuple[dashboard.ThreadingHTTPServer, threading.Thread, str]:
    dashboard.configure_runtime(review_root)
    server = dashboard.ThreadingHTTPServer(("127.0.0.1", 0), dashboard.DashboardHandler)
    thread = threading.Thread(target=server.serve_forever, name="prod006-dashboard", daemon=True)
    thread.start()
    host, port = server.server_address
    return server, thread, f"http://{host}:{port}"


def _stop_dashboard(server: dashboard.ThreadingHTTPServer, thread: threading.Thread) -> None:
    server.shutdown()
    server.server_close()
    thread.join(timeout=10)
    assert not thread.is_alive(), "owned Dashboard server did not stop"


def _write_source_inputs(project: Path) -> None:
    pdf = b"%PDF-1.4\n% PROD-006 tiny non-sensitive source\n"
    pdf_path = project / "00_sources/manual_upload/inbox/prod006-main.pdf"
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    pdf_path.write_bytes(pdf)
    pdf_sha256 = _sha256_bytes(pdf)

    markdown = (
        "# Results\n\n"
        "The tiny source reports a bounded conversion result.\n\n"
        "# References\n\n"
        "[1] Tiny source record.\n"
    )
    slug = "prod006-main"
    extracted = project / f"01_evidence/parses/extracted/{slug}"
    extracted.mkdir(parents=True, exist_ok=True)
    (project / f"01_evidence/mineru/markdown/{slug}.md").parent.mkdir(
        parents=True, exist_ok=True
    )
    (project / f"01_evidence/mineru/markdown/{slug}.md").write_text(
        markdown, encoding="utf-8"
    )
    (extracted / "full.md").write_text(markdown, encoding="utf-8")
    content = [
        {
            "type": "text",
            "page_idx": 0,
            "bbox": [0, 0, 1, 1],
            "text": "The tiny source reports a bounded conversion result.",
        }
    ]
    _write_json(extracted / f"{slug}_content_list.json", content)
    _write_json(extracted / f"{slug}_content_list_v2.json", [content])
    _write_json(extracted / "layout.json", {"pages": [{"page": 1}]})

    reading = b"Tiny source reading layer\n"
    layout = b"Tiny source layout layer\n"
    text_layers = project / "01_evidence/text_layers"
    text_layers.mkdir(parents=True, exist_ok=True)
    reading_path = text_layers / f"{slug}-reading.txt"
    layout_path = text_layers / f"{slug}-layout.json"
    reading_path.write_bytes(reading)
    layout_path.write_bytes(layout)
    _write_json(
        text_layers / "text_layers.manifest.json",
        {
            "sources": [
                {
                    "source_id": SOURCE_ID,
                    "pdf_sha256": pdf_sha256,
                    "page_count": 1,
                    "reading_order_path": reading_path.name,
                    "reading_order_sha256": _sha256_bytes(reading),
                    "layout_path": layout_path.name,
                    "layout_sha256": _sha256_bytes(layout),
                }
            ]
        },
    )
    _write_json(
        project / "01_evidence/mineru/manifest.json",
        {
            "completed": [
                {
                    "relative_pdf_path": "manual_upload/inbox/prod006-main.pdf",
                    "slug": slug,
                    "state": "done",
                }
            ]
        },
    )
    _write_json(
        project / "00_sources/acquisition_final_receipt.json",
        {
            "studies": [
                {
                    "study_id": STUDY_ID,
                    "doi": "10.1000/prod006-tiny",
                    "title": "Tiny non-sensitive source",
                    "main_pdf": {
                        "path": "manual_upload/inbox/prod006-main.pdf",
                        "sha256": pdf_sha256,
                        "size_bytes": len(pdf),
                    },
                }
            ]
        },
    )
    _write_json(
        project / "00_discovery/acquisition_manifest.json",
        {
            "downloads": [
                {
                    "study_id": STUDY_ID,
                    "doi": "10.1000/prod006-tiny",
                    "status": "VERIFIED_EXISTING",
                }
            ]
        },
    )
    _write_json(
        project / "00_sources/source_coverage.json",
        {"studies": [{"study_id": STUDY_ID, "si_policy": "NOT_REQUIRED"}]},
    )
    _write_json(
        project / "00_sources/source_identity_audit.json",
        {
            "results": [
                {
                    "candidate_id": STUDY_ID,
                    "doi": "10.1000/prod006-tiny",
                    "title": "Tiny non-sensitive source",
                }
            ]
        },
    )
    _write_json(
        project / "00_brief/review_state.json",
        {"project_id": PROJECT_ID, "topic": "Tiny source-to-release product path"},
    )


def _approve_parse_quality(project: Path, study_id: str) -> None:
    gate = write_parse_quality_gate(project, study_id)
    for row in gate["objects"]:
        if row["status"] == "usable" and row["review_state"] == "not_required":
            continue
        gate = apply_parse_quality_decision(
            project,
            study_id,
            {
                "object_id": row["object_id"],
                "gate_digest": gate["gate_digest"],
                "object_digest": row["object_digest"],
                "action": "approve_candidate_extraction",
                "note": "Reviewed against the isolated tiny source.",
            },
        )
    state = project_parse_quality_state(project)
    assert state["workflow_can_continue"] is True
    assert state["automatic_extraction_allowed"] is True


def _register_evidence(project: Path) -> dict[str, Any]:
    registered = register_paper_evidence_candidates(
        project,
        STUDY_ID,
        {
            "evidence_id": EVIDENCE_ID,
            "source_id": SOURCE_ID,
            "epistemic_type": "experimental_observation",
            "statement": "The tiny source reports a bounded conversion result.",
            "locator": {
                "source_mode": "parsed_candidate",
                "page": 1,
                "section_or_item": "Results",
                "figure_or_table": None,
                "exact_quote": "The tiny source reports a bounded conversion result.",
            },
            "reported_conditions": ["ambient fixture conditions"],
            "quantitative_results": ["bounded conversion result"],
            "limitations": ["One tiny isolated study; no cross-study comparison."],
            "mechanism_grade": "not_applicable",
            "risk_classes": ["AI_PROVISIONAL", "GAP", "NON_COMPARABLE"],
            # Omitted on input deliberately: the production normalizer binds
            # the persisted empty field_dependencies list without enabling a
            # chemical dual route for this non-chemical evidence row.
        },
    )
    candidate = registered["candidates"][0]
    approved = apply_paper_evidence_decision(
        project,
        {
            "evidence_id": EVIDENCE_ID,
            "candidate_digest": candidate["candidate_digest"],
            "bound_parse_object_digests": candidate["bound_parse_object_digests"],
            "source_pdf_sha256": candidate["source_pdf_sha256"],
            "action": "approve",
            "reason": "Human-reviewed the source-bound tiny observation.",
        },
    )
    assert approved["status"] == "approved"
    state = paper_evidence_state(project)
    assert state["workflow_can_continue"] is True
    row = next(row for row in state["rows"] if row["evidence_id"] == EVIDENCE_ID)
    assert row["field_dependencies"] == []
    assert set(row["risk_classes"]) == {"AI_PROVISIONAL", "GAP", "NON_COMPARABLE"}
    return row


def _register_synthesis_and_contract(project: Path) -> None:
    protocol = register_comparison_protocol(
        project,
        {
            "comparison_id": "comparison-prod006",
            "comparison_objects": [EVIDENCE_ID],
            "axes": ["reported outcome"],
            "normalization_rules": ["Keep source-reported wording and units."],
            "missing_value_policy": "Missing values remain unknown.",
            "incomparability_rules": ["Do not compare absent studies."],
            "counterevidence_rules": ["Record unresolved counterevidence explicitly."],
            "claim_strength": "bounded",
        },
    )
    protocol = apply_comparison_protocol_decision(
        project,
        {
            "comparison_id": protocol["comparison_id"],
            "action": "approve",
            "reason": "Approved the bounded comparison protocol.",
        },
    )
    assert protocol["decision"]["action"] == "approve"
    claim = register_synthesis_candidates(
        project,
        {
            "synthesis_id": SYNTHESIS_ID,
            "proposition": "The source reports a bounded outcome.",
            "comparison_axis": "reported outcome",
            "supporting_evidence_ids": [EVIDENCE_ID],
            "counter_evidence_ids": [],
            "applicability_boundary": "Only this isolated tiny source.",
            "mechanism_evidence_grade": "not_applicable",
            "uncertainty": "Single study; non-comparable to absent studies.",
            "risk_class": "NON_COMPARABLE",
            "single_study": True,
        },
    )["claims"][0]
    apply_synthesis_decision(
        project,
        {
            "synthesis_id": claim["synthesis_id"],
            "action": "approve",
            "reason": "Approved bounded wording with the stated limitation.",
        },
    )
    synthesis = synthesis_state(project)
    assert synthesis["workflow_can_continue"] is True

    contract = register_section_contracts(
        project,
        {
            "section_id": SECTION_ID,
            "research_question": "What was reported by the isolated source?",
            "comparison_axes": ["reported outcome"],
            "expected_synthesis": "Retain bounded source wording.",
            "counterevidence_and_limitations": [
                "One source only; cross-study comparison is not available."
            ],
            "evidence_budget": 1,
            "synthesis_budget": 1,
            "figure_plan": [{"kind": "source", "purpose": "show source-bound fixture figure"}],
            "allowed_wording_strength": "bounded",
        },
    )["contracts"][0]
    apply_section_contract_decision(
        project,
        {
            "section_id": contract["section_id"],
            "action": "approve",
            "reason": "Approved the source-bound section contract.",
        },
    )
    contracts = section_contract_state(project)
    assert contracts["workflow_can_continue"] is True


def _write_source_figure_registry(project: Path) -> dict[str, Any]:
    image_path = project / "01_evidence/parses/extracted/prod006-main/images/figure1.png"
    image_path.parent.mkdir(parents=True, exist_ok=True)
    image_path.write_bytes(_TINY_PNG)
    image_sha256 = _sha256_bytes(_TINY_PNG)
    figure = {
        "figure_id": f"{STUDY_ID}:{SOURCE_ID}:figure-1",
        "study_id": STUDY_ID,
        "source_id": SOURCE_ID,
        "page": 1,
        "figure_label": "Figure 1",
        "caption": "Tiny source figure",
        "asset_path": "01_evidence/parses/extracted/prod006-main/images/figure1.png",
        "asset_sha256": image_sha256,
        "source_pdf_sha256": load_source_truth_bundle(project, STUDY_ID)["sources"][0]["pdf"]["sha256"],
        "evidence_ids": [EVIDENCE_ID],
        "selection_status": "selected",
        "fragments": [
            {
                "page": 1,
                "block_index": 0,
                "bbox": [0, 0, 1, 1],
                "asset_path": "01_evidence/parses/extracted/prod006-main/images/figure1.png",
                "asset_sha256": image_sha256,
                "caption_association": "explicit_caption_anchor",
            }
        ],
        "selection_reason": "Tiny source figure selected for source-bound release verification.",
    }
    source_truth_digest = _source_truth_digest(project)
    content_list_v2_digest = _content_list_v2_digest(project)
    body = {
        "source_truth_digest": source_truth_digest,
        "content_list_v2_digest": content_list_v2_digest,
        "chemical_paper_project_binding_digest": None,
        "figures": [figure],
        "locator_gaps": [],
    }
    registry = {
        "schema_version": "review-writer-source-figure-registry.v1",
        "project_id": project.name,
        "figure_policy": "source_figures_or_synthesis_placeholders_only",
        "figures": [figure],
        "locator_gaps": [],
        **body,
        "registry_digest": canonical_digest(body),
    }
    _write_json(project / "03_figures/source_figure_registry.json", registry)
    return figure


def _build_project(review_root: Path) -> Path:
    project = review_root / PROJECT_ID
    project.mkdir(parents=True, exist_ok=True)
    _write_source_inputs(project)
    write_source_truth_bundle(project, STUDY_ID)
    _approve_parse_quality(project, STUDY_ID)
    _register_evidence(project)
    _register_synthesis_and_contract(project)
    figure = _write_source_figure_registry(project)
    attribution = (
        f"Source Figure Attribution: {figure['figure_id']} | {SOURCE_ID} | page 1 | Figure 1"
    )
    body = (
        "[synthesis:synthesis-prod006-bounded] The source reports a bounded outcome.\n\n"
        "![Tiny source figure](../01_evidence/parses/extracted/prod006-main/images/figure1.png)\n\n"
        f"{attribution} [evidence:{EVIDENCE_ID}]"
    )
    draft = register_section_draft(
        project,
        {
            "section_id": SECTION_ID,
            "heading": "Reported result",
            "body": body,
            "generation_content_agent_result_digest": "a" * 64,
        },
    )
    assert draft["status"] == "needs_human_edit"
    edited_body = body.replace("reports a bounded outcome", "records a bounded outcome")
    approved = approve_section(
        project,
        SECTION_ID,
        {"actor_type": "human_researcher", "actor_label": "prod006-human"},
        edited_body=edited_body,
        reason="Human reviewed the provisional, gap, and non-comparable wording.",
        expected_draft_digest=draft["draft_digest"],
    )
    assert approved["decision"]["action"] == "approve"
    merged = merge_authoritative_manuscript(project)
    assert merged["status"] == "approved"
    assert (project / "04_manuscript/manuscript.md").is_file()
    return project


def _assert_release_payload(project: Path, payload: dict[str, Any]) -> dict[str, bytes]:
    release_dir = project / "05_release"
    markdown_path = release_dir / "self_reviewed_draft.md"
    docx_path = release_dir / "self_reviewed_draft.docx"
    snapshot_path = release_dir / "release_snapshot.json"
    quality_path = release_dir / "quality_report.json"
    paths = (markdown_path, docx_path, snapshot_path, quality_path)
    assert all(path.is_file() and not path.is_symlink() for path in paths)
    snapshot = json.loads(snapshot_path.read_text(encoding="utf-8"))
    quality = json.loads(quality_path.read_text(encoding="utf-8"))
    for filename, value in (
        ("release_snapshot.v1.schema.json", snapshot),
        ("project_release.v2.schema.json", quality),
    ):
        schema = json.loads((Path(__file__).parents[2] / "schemas/delivery" / filename).read_text(encoding="utf-8"))
        errors = list(Draft202012Validator(schema).iter_errors(value))
        assert not errors, f"{filename}: {errors}"

    authoritative = project / "04_manuscript/manuscript.md"
    markdown_bytes = markdown_path.read_bytes()
    docx_bytes = docx_path.read_bytes()
    assert markdown_bytes == authoritative.read_bytes()
    assert snapshot["schema_version"] == "release-snapshot.v1"
    assert snapshot["release_level"] == "SELF_REVIEWED_DRAFT"
    assert snapshot["status"] == "SELF_REVIEWED_DRAFT"
    assert snapshot["markdown_path"] == "05_release/self_reviewed_draft.md"
    assert snapshot["docx_path"] == "05_release/self_reviewed_draft.docx"
    assert snapshot["manuscript_sha256"] == _sha256_bytes(authoritative.read_bytes())
    assert snapshot["release_markdown_sha256"] == _sha256_bytes(markdown_bytes)
    assert snapshot["docx_sha256"] == _sha256_bytes(docx_bytes)
    assert quality["schema_version"] == "project-release.v2"
    assert quality["release_level"] == "SELF_REVIEWED_DRAFT"
    assert quality["manuscript_sha256"] == snapshot["manuscript_sha256"]
    assert quality["release_markdown_sha256"] == snapshot["release_markdown_sha256"]
    assert quality["docx_sha256"] == snapshot["docx_sha256"]
    assert snapshot["integrity"]["markdown_roundtrip_match"] is True
    assert snapshot["integrity"]["legacy_repackage_only"] is False
    assert dashboard.new_route_release_docx_is_current(docx_path) is True

    document = Document(docx_path)
    document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "records a bounded outcome" in document_text
    assert "Source Figure Attribution" in document_text
    assert payload["release_level"] == "SELF_REVIEWED_DRAFT"
    return {str(path.relative_to(project)): path.read_bytes() for path in paths}


def _assert_views(base_url: str, project: Path, *, release_current: bool = False) -> None:
    for path, marker in (
        ("/review", "证据与综合判断工作台"),
        ("/draft", "Draft"),
        ("/final", "Release Report"),
    ):
        status, html = _http_text(base_url, path)
        assert status == 200
        assert marker in html

    source_status, source_bytes = _http_bytes(
        base_url,
        f"/api/project/{PROJECT_ID}/source/{SOURCE_ID}/pdf",
    )
    assert source_status == 200
    assert source_bytes.startswith(b"%PDF-1.4")

    evidence_status, evidence = _http_json(
        base_url, f"/api/project/{PROJECT_ID}/paper-evidence"
    )
    assert evidence_status == 200
    evidence_rows = evidence.get("items", [])
    row = next(row for row in evidence_rows if row["evidence_id"] == EVIDENCE_ID)
    assert set(row["risk_classes"]) == {"AI_PROVISIONAL", "GAP", "NON_COMPARABLE"}
    persisted_row = next(
        row for row in paper_evidence_state(project)["rows"] if row["evidence_id"] == EVIDENCE_ID
    )
    assert persisted_row["field_dependencies"] == []

    draft_status, draft = _http_json(base_url, f"/api/project/{PROJECT_ID}/draft")
    assert draft_status == 200
    assert draft["available"] is True
    assert any(section["section_id"] == SECTION_ID for section in draft["sections"])

    final_status, final = _http_json(base_url, f"/api/project/{PROJECT_ID}/final")
    assert final_status == 200
    assert final["manuscript_source"] == (
        "release_snapshot" if release_current else "authoritative_manuscript"
    )
    if release_current:
        assert final["release_snapshot"]["matches_authoritative"] is True
        assert final["release_snapshot"]["docx_exists"] is True
    assert "records a bounded outcome" in final["final_draft_md"]


def test_prod006_product_use_closes_source_to_release_and_survives_cold_restart() -> None:
    with tempfile.TemporaryDirectory(prefix="prod006-product-use-") as temporary_root:
        review_root = Path(temporary_root)
        project = _build_project(review_root)
        server, thread, base_url = _start_dashboard(review_root)
        try:
            _assert_views(base_url, project)
            status, result = _post_release(base_url)
            assert status == 200
            assert result["ok"] is True
            first_release_bytes = _assert_release_payload(project, result)
            _assert_views(base_url, project, release_current=True)
        finally:
            _stop_dashboard(server, thread)

        # Cold restart the same isolated project and verify that current is a
        # persisted binding, not an in-memory server result.
        server, thread, base_url = _start_dashboard(review_root)
        try:
            _assert_views(base_url, project, release_current=True)
            final_status, final = _http_json(base_url, f"/api/project/{PROJECT_ID}/final")
            assert final_status == 200
            assert final["release_snapshot"]["matches_authoritative"] is True
            current_release_bytes = _assert_release_payload(
                project,
                {
                    "release_level": final["release_status"],
                },
            )
            assert current_release_bytes == first_release_bytes
        finally:
            _stop_dashboard(server, thread)
