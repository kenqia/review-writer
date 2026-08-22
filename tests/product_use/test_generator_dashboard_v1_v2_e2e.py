from __future__ import annotations

import hashlib
import json
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any
from urllib.error import HTTPError
from urllib.parse import quote
from urllib.request import Request, urlopen

from docx import Document

from review_writer.product_foundation import VersionContext
from review_writer.product_foundation.project_root import version_context_root
from review_writer.project.source_truth import write_source_truth_bundle
from tests.product_use import test_prod006_source_to_release as prod006_fixture


PROJECT_ID = prod006_fixture.PROJECT_ID
STUDY_ID = prod006_fixture.STUDY_ID
EVIDENCE_ID = prod006_fixture.EVIDENCE_ID
REPO_ROOT = Path(__file__).resolve().parents[2]
CLI = REPO_ROOT / "scripts" / "run_vertical_review.py"


def _json_request(
    base_url: str,
    path: str,
    *,
    method: str = "GET",
    payload: dict[str, Any] | None = None,
) -> tuple[int, dict[str, Any]]:
    body = None
    headers: dict[str, str] = {}
    if payload is not None:
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        headers = {
            "Content-Type": "application/json",
            "Content-Length": str(len(body)),
        }
    request = Request(f"{base_url}{path}", data=body, headers=headers, method=method)
    try:
        with urlopen(request, timeout=15) as response:
            return response.status, json.loads(response.read().decode("utf-8"))
    except HTTPError as exc:
        return exc.code, json.loads(exc.read().decode("utf-8"))


def _raw_request(
    base_url: str,
    path: str,
    *,
    method: str = "GET",
    payload: dict[str, Any] | None = None,
) -> tuple[int, bytes, dict[str, str]]:
    body = (
        json.dumps(payload, ensure_ascii=False).encode("utf-8")
        if payload is not None
        else None
    )
    headers = (
        {
            "Content-Type": "application/json",
            "Content-Length": str(len(body)),
        }
        if body is not None
        else {}
    )
    request = Request(f"{base_url}{path}", data=body, headers=headers, method=method)
    try:
        with urlopen(request, timeout=30) as response:
            return response.status, response.read(), dict(response.headers.items())
    except HTTPError as exc:
        return exc.code, exc.read(), dict(exc.headers.items())


def _build_project(review_root: Path, *, with_source_figure: bool = False) -> Path:
    project = review_root / PROJECT_ID
    project.mkdir(parents=True, exist_ok=True)
    prod006_fixture._write_source_inputs(project)
    if with_source_figure:
        _add_synthetic_source_figure(project, rebuild_source_truth=False)
    write_source_truth_bundle(project, STUDY_ID)
    prod006_fixture._approve_parse_quality(project, STUDY_ID)
    prod006_fixture._register_evidence(project)
    return project


def _add_synthetic_source_figure(
    project: Path, *, rebuild_source_truth: bool = True
) -> None:
    """Add an explicitly synthetic, source-bound image before rebuilding Source Truth."""
    extracted = project / "01_evidence/parses/extracted/prod006-main"
    image_path = extracted / "images/synthetic_figure.png"
    image_path.parent.mkdir(parents=True, exist_ok=True)
    image_path.write_bytes(prod006_fixture._TINY_PNG)
    content_path = extracted / "prod006-main_content_list_v2.json"
    content_path.write_text(
        json.dumps(
            [
                [
                    {
                        "type": "image",
                        "page_idx": 0,
                        "bbox": [0, 0, 100, 100],
                        "content": {
                            "image_source": {"path": "images/synthetic_figure.png"},
                            "image_caption": [
                                {
                                    "content": (
                                        "Figure 1. Synthetic non-sensitive fixture image; "
                                        "not a scientific result."
                                    )
                                }
                            ],
                        },
                    }
                ]
            ],
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )
    if rebuild_source_truth:
        write_source_truth_bundle(project, STUDY_ID)


def _protocol_payload(
    *,
    comparison_id: str = "comparison-generator-dashboard",
    evidence_digest: str | None = None,
) -> dict[str, Any]:
    payload: dict[str, Any] = {
        "comparison_id": comparison_id,
        "comparison_objects": [EVIDENCE_ID],
        "axes": ["reported outcome"],
        "normalization_rules": ["Keep source-reported wording and units."],
        "missing_value_policy": "Missing values remain unknown.",
        "incomparability_rules": ["Do not compare absent studies."],
        "counterevidence_rules": ["Record unresolved counterevidence explicitly."],
        "claim_strength": "bounded",
    }
    if evidence_digest is not None:
        payload["paper_evidence_projection_digest"] = evidence_digest
    return payload


def _figure_placeholder_payload(
    *,
    placeholder_id: str = "figure-generator-bounded",
    synthesis_id: str = "synthesis-generator-bounded",
) -> dict[str, Any]:
    return {
        "placeholder_id": placeholder_id,
        "scientific_question": "How should this bounded outcome be compared without overclaiming?",
        "reader_takeaway": "The comparison remains bounded until a human figure is supplied.",
        "panels": [
            {
                "panel": "A",
                "task": "Show the source-bound outcome and retain the single-study limitation.",
                "synthesis_claim_ids": [synthesis_id],
                "source_figure_ids": [],
            }
        ],
        "comparison_axis": "reported outcome",
        "required_labels_units": ["source-reported outcome"],
        "counter_evidence": ["No cross-study comparator is available."],
        "forbidden_overclaims": ["Do not imply generalizability beyond this source."],
        "unresolved_uncertainties": ["A human must supply and verify the synthesis figure."],
        "caption_draft": "Bounded source outcome with an explicit human-figure gap.",
        "target_size": "single-column",
        "status": "awaiting_human_figure",
    }


def _run_cli(*arguments: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(CLI), *arguments],
        cwd=REPO_ROOT,
        capture_output=True,
        check=False,
        text=True,
    )


def _cli_json(result: subprocess.CompletedProcess[str]) -> dict[str, Any]:
    stream = result.stdout if result.returncode == 0 else result.stderr
    return json.loads(stream)


def _snapshot_files(root: Path) -> dict[str, bytes]:
    return {
        path.relative_to(root).as_posix(): path.read_bytes()
        for path in sorted(root.rglob("*"), key=lambda item: item.as_posix())
        if path.is_file() and not path.is_symlink()
    }


def test_public_comparison_protocol_approve_is_not_blocked_on_fresh_project() -> None:
    """The public approval seam must work after Agent-side upstream preparation."""
    with tempfile.TemporaryDirectory(prefix="generator-dashboard-v1-v2-") as temporary_root:
        review_root = Path(temporary_root)
        project = _build_project(review_root)
        protocol_input = review_root / "comparison-protocol.json"
        protocol_input.write_text(
            json.dumps(_protocol_payload(), ensure_ascii=False) + "\n",
            encoding="utf-8",
        )

        server, thread, base_url = prod006_fixture._start_dashboard(review_root)
        try:
            status, before = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/comparison-protocol",
            )
            assert status == 200
            assert before["evidence_ready"] is True
            protocol = before["protocol"]
            assert protocol["version_token"]

            # This is the frozen first public failure before the Agent
            # producer runs: Dashboard is intentionally decision-only.
            status, after = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/comparison-protocol",
                method="PUT",
                payload={
                    "action": "approve",
                    "reason": "Human approved the bounded comparison protocol.",
                    "version_token": protocol["version_token"],
                    "actor_type": "human_researcher",
                    "actor_label": "generator-dashboard-human",
                },
            )
            assert status == 400
            assert after == {"error": "决定未保存，请检查后重试"}
        finally:
            prod006_fixture._stop_dashboard(server, thread)

        produced = _run_cli(
            "register-comparison-protocol",
            "--project",
            str(project),
            "--input",
            str(protocol_input),
        )
        assert produced.returncode == 0, produced.stderr
        assert _cli_json(produced)["status"] == "REGISTERED"

        server, thread, base_url = prod006_fixture._start_dashboard(review_root)
        try:
            status, before = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/comparison-protocol",
            )
            assert status == 200
            protocol = before["protocol"]
            assert protocol["comparison_objects"] == [EVIDENCE_ID]

            status, after = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/comparison-protocol",
                method="PUT",
                payload={
                    "action": "approve",
                    "reason": "Human approved the bounded comparison protocol.",
                    "version_token": protocol["version_token"],
                    "actor_type": "human_researcher",
                    "actor_label": "generator-dashboard-human",
                },
            )
            assert status == 200, after
            assert after["workflow_can_continue"] is True
            assert after["protocol"]["decision"]["action"] == "approve"
        finally:
            prod006_fixture._stop_dashboard(server, thread)


def test_agent_protocol_producer_rejects_invalid_stale_and_duplicate_inputs_without_write() -> None:
    cases = (
        ("invalid", {"comparison_id": "invalid"}, "COMPARISON_PROTOCOL_INVALID"),
        (
            "stale",
            _protocol_payload(evidence_digest="0" * 64),
            "COMPARISON_PROTOCOL_STALE",
        ),
        (
            "duplicate",
            {**_protocol_payload(), "comparison_objects": [EVIDENCE_ID, EVIDENCE_ID]},
            "COMPARISON_PROTOCOL_DUPLICATE",
        ),
    )
    for name, payload, expected_code in cases:
        with tempfile.TemporaryDirectory(prefix=f"generator-protocol-{name}-") as temporary_root:
            review_root = Path(temporary_root)
            project = _build_project(review_root)
            input_path = review_root / f"{name}.json"
            input_path.write_text(
                json.dumps(payload, ensure_ascii=False) + "\n",
                encoding="utf-8",
            )
            baseline = _snapshot_files(project)

            result = _run_cli(
                "register-comparison-protocol",
                "--project",
                str(project),
                "--input",
                str(input_path),
            )
            assert result.returncode == 2
            assert _cli_json(result) == {
                "command": "register-comparison-protocol",
                "error_code": expected_code,
                "status": "ERROR",
            }, name
            assert _snapshot_files(project) == baseline, name
            assert not (project / "02_synthesis/comparison_protocol.json").exists(), name


def test_agent_figure_placeholder_producer_is_visible_through_dashboard() -> None:
    """Agent figure briefs must use the canonical placeholder owner and remain reviewable."""
    with tempfile.TemporaryDirectory(prefix="generator-figure-placeholder-") as temporary_root:
        review_root = Path(temporary_root)
        project = _build_project(review_root)
        input_path = review_root / "figure-placeholder.json"
        input_path.write_text(
            json.dumps(_figure_placeholder_payload(), ensure_ascii=False) + "\n",
            encoding="utf-8",
        )

        result = _run_cli(
            "register-synthesis-figure-placeholder",
            "--project",
            str(project),
            "--input",
            str(input_path),
        )
        assert result.returncode == 0, result.stderr
        summary = _cli_json(result)
        assert summary == {
            "command": "register-synthesis-figure-placeholder",
            "placeholder_id": "figure-generator-bounded",
            "status": "REGISTERED",
        }

        server, thread, base_url = prod006_fixture._start_dashboard(review_root)
        try:
            status, figures = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/review-figures",
            )
            assert status == 200
            assert figures["summary"]["placeholder_count"] == 1
            placeholder = figures["placeholders"][0]
            assert placeholder["placeholder_id"] == "figure-generator-bounded"
            assert placeholder["panels"][0]["synthesis_claim_ids"] == [
                "synthesis-generator-bounded"
            ]
            assert placeholder["status"] == "awaiting_human_figure"
        finally:
            prod006_fixture._stop_dashboard(server, thread)


def test_agent_source_figure_registry_producer_is_canonical_and_source_bound() -> None:
    """The Agent producer must invoke the existing Source Figure registry owner."""
    with tempfile.TemporaryDirectory(prefix="generator-source-figure-registry-") as temporary_root:
        review_root = Path(temporary_root)
        project = _build_project(review_root)
        _add_synthetic_source_figure(project)

        result = _run_cli(
            "build-source-figure-registry",
            "--project",
            str(project),
        )
        assert result.returncode == 0, result.stderr
        summary = _cli_json(result)
        assert summary["command"] == "build-source-figure-registry"
        assert summary["status"] == "REGISTERED"
        registry_path = project / "03_figures/source_figure_registry.json"
        registry = json.loads(registry_path.read_text(encoding="utf-8"))
        assert registry["figure_policy"] == "source_figures_or_synthesis_placeholders_only"
        assert registry["selected_count"] == 1
        figure = registry["figures"][0]
        assert figure["source_id"] == prod006_fixture.SOURCE_ID
        assert figure["asset_path"].endswith("images/synthetic_figure.png")
        assert "Synthetic non-sensitive fixture image" in figure["caption"]


def test_agent_figure_placeholder_producer_rejects_invalid_and_conflicting_inputs_without_write() -> None:
    cases = (
        ("invalid", {"placeholder_id": "invalid"}, "PLACEHOLDER_INVALID"),
        (
            "conflict",
            {
                **_figure_placeholder_payload(),
                "scientific_question": "Conflicting figure brief.",
            },
            "PLACEHOLDER_CONFLICT",
        ),
    )
    for name, payload, expected_code in cases:
        with tempfile.TemporaryDirectory(prefix=f"generator-figure-placeholder-{name}-") as temporary_root:
            review_root = Path(temporary_root)
            project = _build_project(review_root)
            valid_input = review_root / "valid-figure-placeholder.json"
            valid_input.write_text(
                json.dumps(_figure_placeholder_payload(), ensure_ascii=False) + "\n",
                encoding="utf-8",
            )
            initial = _run_cli(
                "register-synthesis-figure-placeholder",
                "--project",
                str(project),
                "--input",
                str(valid_input),
            )
            assert initial.returncode == 0, initial.stderr

            input_path = review_root / f"{name}-figure-placeholder.json"
            input_path.write_text(
                json.dumps(payload, ensure_ascii=False) + "\n",
                encoding="utf-8",
            )
            baseline = _snapshot_files(project)
            result = _run_cli(
                "register-synthesis-figure-placeholder",
                "--project",
                str(project),
                "--input",
                str(input_path),
            )
            assert result.returncode == 2
            assert _cli_json(result) == {
                "command": "register-synthesis-figure-placeholder",
                "error_code": expected_code,
                "status": "ERROR",
            }, name
            assert _snapshot_files(project) == baseline, name


def test_agent_synthesis_candidates_follow_public_protocol_approval() -> None:
    """The Agent must register current synthesis before Dashboard can review it."""
    comparison_id = "comparison-generator-synthesis"
    synthesis_id = "synthesis-generator-bounded"
    with tempfile.TemporaryDirectory(prefix="generator-synthesis-") as temporary_root:
        review_root = Path(temporary_root)
        project = _build_project(review_root, with_source_figure=True)
        protocol_input = review_root / "comparison-protocol.json"
        protocol_input.write_text(
            json.dumps(
                _protocol_payload(comparison_id=comparison_id),
                ensure_ascii=False,
            )
            + "\n",
            encoding="utf-8",
        )
        protocol_result = _run_cli(
            "register-comparison-protocol",
            "--project",
            str(project),
            "--input",
            str(protocol_input),
        )
        assert protocol_result.returncode == 0, protocol_result.stderr

        server, thread, base_url = prod006_fixture._start_dashboard(review_root)
        try:
            status, protocol = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/comparison-protocol",
            )
            assert status == 200
            protocol_token = protocol["protocol"]["version_token"]
            status, approved_protocol = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/comparison-protocol",
                method="PUT",
                payload={
                    "action": "approve",
                    "reason": "Approved the bounded protocol before synthesis.",
                    "version_token": protocol_token,
                    "actor_type": "human_researcher",
                    "actor_label": "generator-synthesis-human",
                },
            )
            assert status == 200
            assert approved_protocol["workflow_can_continue"] is True
        finally:
            prod006_fixture._stop_dashboard(server, thread)

        coverage_input = review_root / "coverage-map.json"
        coverage_input.write_text(
            json.dumps(
                {
                    "comparison_id": comparison_id,
                    "axes": [
                        {
                            "axis_id": "reported outcome",
                            "counterevidence_ids": [],
                            "incomparable_items": [],
                        }
                    ],
                    "known_omissions": ["One synthetic study; no cross-study comparison."],
                },
                ensure_ascii=False,
            )
            + "\n",
            encoding="utf-8",
        )
        coverage_result = _run_cli(
            "register-coverage-map",
            "--project",
            str(project),
            "--input",
            str(coverage_input),
        )
        assert coverage_result.returncode == 0, coverage_result.stderr

        synthesis_input = review_root / "synthesis.json"
        synthesis_input.write_text(
            json.dumps(
                {
                    "synthesis_id": synthesis_id,
                    "proposition": "The synthetic source reports a bounded outcome.",
                    "comparison_axis": "reported outcome",
                    "supporting_evidence_ids": [EVIDENCE_ID],
                    "counter_evidence_ids": [],
                    "applicability_boundary": "Only this isolated synthetic source.",
                    "mechanism_evidence_grade": "not_applicable",
                    "uncertainty": "Single study; cross-study comparison is unavailable.",
                    "risk_class": "NON_COMPARABLE",
                    "single_study": True,
                },
                ensure_ascii=False,
            )
            + "\n",
            encoding="utf-8",
        )
        synthesis_result = _run_cli(
            "register-synthesis-candidates",
            "--project",
            str(project),
            "--input",
            str(synthesis_input),
        )
        assert synthesis_result.returncode == 0, synthesis_result.stderr

        server, thread, base_url = prod006_fixture._start_dashboard(review_root)
        try:
            status, synthesis = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/synthesis",
            )
            assert status == 200
            item = next(row for row in synthesis["items"] if row["synthesis_id"] == synthesis_id)
            assert item["status"] == "needs_review"

            status, approved = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/synthesis",
                method="PUT",
                payload={
                    "synthesis_id": synthesis_id,
                    "action": "approve",
                    "reason": "Human approved bounded synthesis with its limitation.",
                    "version_token": item["version_token"],
                    "actor_type": "human_researcher",
                    "actor_label": "generator-synthesis-human",
                },
            )
            assert status == 200, approved
            assert approved["workflow_can_continue"] is True
            assert approved["items"][0]["status"] == "approved"

            contract_input = review_root / "section-contract.json"
            contract_input.write_text(
                json.dumps(
                    {
                        "section_id": "reported-result",
                        "research_question": "Which source evidence supports the synthetic outcome?",
                        "comparison_axes": ["reported outcome"],
                        "expected_synthesis": "Retain bounded source wording.",
                        "counterevidence_and_limitations": [
                            "One synthetic source; cross-study comparison is unavailable."
                        ],
                        "evidence_budget": 1,
                        "synthesis_budget": 1,
                        "figure_plan": [
                            {"kind": "source", "purpose": "show source-bound figure"}
                        ],
                        "allowed_wording_strength": "bounded",
                    },
                    ensure_ascii=False,
                )
                + "\n",
                encoding="utf-8",
            )
            contract_result = _run_cli(
                "register-section-contracts",
                "--project",
                str(project),
                "--input",
                str(contract_input),
            )
            assert contract_result.returncode == 0, contract_result.stderr

            status, contracts = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/section-contracts",
            )
            assert status == 200
            contract = next(
                row for row in contracts["items"] if row["section_id"] == "reported-result"
            )
            assert contract["status"] == "needs_review"
            status, approved_contract = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/section-contracts",
                method="PUT",
                payload={
                    "section_id": "reported-result",
                    "action": "approve",
                    "reason": "Human approved the bounded section contract.",
                    "version_token": contract["version_token"],
                    "actor_type": "human_researcher",
                    "actor_label": "generator-synthesis-human",
                },
            )
            assert status == 200, approved_contract
            assert approved_contract["workflow_can_continue"] is True
        finally:
            prod006_fixture._stop_dashboard(server, thread)

        draft_body = (
            f"[synthesis:{synthesis_id}] The synthetic source reports a bounded outcome.\n\n"
            f"[evidence:{EVIDENCE_ID}] The isolated source reports a bounded outcome."
        )
        draft_input = review_root / "section-draft.json"
        draft_input.write_text(
            json.dumps(
                {
                    "section_id": "reported-result",
                    "heading": "Reported result",
                    "body": draft_body,
                    "content_agent_result_digest": "d" * 64,
                },
                ensure_ascii=False,
            )
            + "\n",
            encoding="utf-8",
        )
        draft_result = _run_cli(
            "register-section-draft",
            "--project",
            str(project),
            "--input",
            str(draft_input),
        )
        assert draft_result.returncode == 0, draft_result.stderr
        draft_summary = _cli_json(draft_result)
        assert draft_summary["command"] == "register-section-draft"
        assert draft_summary["status"] == "REGISTERED"

        server, thread, base_url = prod006_fixture._start_dashboard(review_root)
        try:
            status, draft = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/draft",
            )
            assert status == 200
            assert draft["available"] is True
            section = next(
                row for row in draft["sections"] if row["section_id"] == "reported-result"
            )
            assert section["status"] == "needs_human_edit"
            assert section["body"] == draft_body
            assert section["claim_bindings"] == [
                {"paper_evidence_ids": [], "synthesis_ids": [synthesis_id]},
                {"paper_evidence_ids": [EVIDENCE_ID], "synthesis_ids": []},
            ]

            edited_body = draft_body.replace(
                "reports a bounded outcome", "records a bounded outcome"
            )
            status, approved_draft = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/draft",
                method="PUT",
                payload={
                    "section_id": section["section_id"],
                    "edited_body": edited_body,
                    "reason": "Human preserved the evidence and synthesis markers while editing wording.",
                    "version_token": section["version_token"],
                    "actor_type": "simulated_researcher_agent",
                    "actor_label": "generator-synthesis-human",
                },
            )
            assert status == 200, approved_draft
            assert approved_draft["status"] == "approved"
            approved_section = next(
                row
                for row in approved_draft["sections"]
                if row["section_id"] == "reported-result"
            )
            assert approved_section["body"] == edited_body
            assert approved_section["claim_bindings"] == section["claim_bindings"]
            assert approved_section["decision"]["action"] == "approve"
            v1_body = approved_section["body"]
            v1_version_token = approved_section["version_token"]
        finally:
            prod006_fixture._stop_dashboard(server, thread)

        v2_addition = (
            f"[synthesis:{synthesis_id}] The limitation remains explicit for this isolated source."
        )
        v2_input = review_root / "section-draft-v2.json"
        v2_input.write_text(
            json.dumps(
                {
                    "section_id": "reported-result",
                    "body": v2_addition,
                    "content_agent_result_digest": "f" * 64,
                },
                ensure_ascii=False,
            )
            + "\n",
            encoding="utf-8",
        )
        v2_result = _run_cli(
            "generate-section-draft-v2",
            "--project",
            str(project),
            "--input",
            str(v2_input),
        )
        assert v2_result.returncode == 0, v2_result.stderr
        v2_summary = _cli_json(v2_result)
        assert v2_summary["command"] == "generate-section-draft-v2"
        assert v2_summary["status"] == "REGISTERED"

        server, thread, base_url = prod006_fixture._start_dashboard(review_root)
        try:
            status, v2_draft = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/draft",
            )
            assert status == 200
            v2_section = next(
                row
                for row in v2_draft["sections"]
                if row["section_id"] == "reported-result"
            )
            assert v2_section["body"] == f"{v1_body}\n\n{v2_addition}"
            assert v2_section["body"].startswith(v1_body)
            assert v2_section["version_token"] != v1_version_token
            assert v2_section["status"] == "needs_human_edit"
            assert v2_section["decision"] is None
            assert v2_section["claim_bindings"] == [
                {"paper_evidence_ids": [], "synthesis_ids": [synthesis_id]},
                {"paper_evidence_ids": [EVIDENCE_ID], "synthesis_ids": []},
                {"paper_evidence_ids": [], "synthesis_ids": [synthesis_id]},
            ]

            v2_edited_body = v2_section["body"].replace(
                "The limitation remains explicit",
                "The limitation stays explicit",
            )
            status, approved_v2 = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/draft",
                method="PUT",
                payload={
                    "section_id": v2_section["section_id"],
                    "edited_body": v2_edited_body,
                    "reason": "Human reviewed the v2 addition and preserved all source markers.",
                    "version_token": v2_section["version_token"],
                    "actor_type": "simulated_researcher_agent",
                    "actor_label": "generator-synthesis-human-v2",
                },
            )
            assert status == 200, approved_v2
            assert approved_v2["status"] == "approved"
            approved_v2_section = next(
                row
                for row in approved_v2["sections"]
                if row["section_id"] == "reported-result"
            )
            assert approved_v2_section["body"] == v2_edited_body
            assert approved_v2_section["decision"]["action"] == "approve"
            assert len(approved_v2_section["claim_bindings"]) == 3

            source_registry_result = _run_cli(
                "build-source-figure-registry",
                "--project",
                str(project),
            )
            assert source_registry_result.returncode == 0, source_registry_result.stderr
            assert _cli_json(source_registry_result)["status"] == "REGISTERED"
            review_figures_status, review_figures = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/review-figures",
            )
            assert review_figures_status == 200, review_figures
            assert review_figures["status"] == "current"
            assert review_figures["summary"]["source_count"] == 1
            source_figure = review_figures["source_figures"][0]
            assert source_figure["selection_status"] == "selected"

            # Exercise the existing public human selection seam explicitly:
            # clear the producer's initial candidate selection, then select
            # the same source figure using the fresh Dashboard token.
            status, deselected = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/review-figures",
                method="PUT",
                payload={
                    "figure_id": source_figure["figure_id"],
                    "selection_status": "available",
                    "version_token": source_figure["version_token"],
                },
            )
            assert status == 200, deselected
            available_figure = deselected["source_figures"][0]
            assert available_figure["selection_status"] == "available"
            status, refreshed_figures = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/review-figures",
            )
            assert status == 200
            available_figure = refreshed_figures["source_figures"][0]
            status, selected_figures = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/review-figures",
                method="PUT",
                payload={
                    "figure_id": available_figure["figure_id"],
                    "selection_status": "selected",
                    "version_token": available_figure["version_token"],
                },
            )
            assert status == 200, selected_figures
            selected_source_figure = selected_figures["source_figures"][0]
            assert selected_source_figure["selection_status"] == "selected"

            placeholder_payload = _figure_placeholder_payload(synthesis_id=synthesis_id)
            placeholder_input = review_root / "figure-placeholder.json"
            placeholder_input.write_text(
                json.dumps(
                    placeholder_payload,
                    ensure_ascii=False,
                )
                + "\n",
                encoding="utf-8",
            )
            placeholder_result = _run_cli(
                "register-synthesis-figure-placeholder",
                "--project",
                str(project),
                "--input",
                str(placeholder_input),
            )
            assert placeholder_result.returncode == 0, placeholder_result.stderr
            placeholder_summary = _cli_json(placeholder_result)
            assert placeholder_summary["status"] == "REGISTERED"
            review_figures_status, review_figures = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/review-figures",
            )
            assert review_figures_status == 200, review_figures
            assert review_figures["summary"]["placeholder_count"] == 1

            # The placeholder is a new upstream dependency.  Reuse the
            # existing Dashboard draft seam to make its brief visible and
            # republish the same authoritative manuscript lineage.
            draft_after_placeholder_status, draft_after_placeholder = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/draft",
            )
            assert draft_after_placeholder_status == 200, draft_after_placeholder
            section_after_placeholder = next(
                row
                for row in draft_after_placeholder["sections"]
                if row["section_id"] == "reported-result"
            )
            visible_placeholder = (
                f"[synthesis:{synthesis_id}] SYNTHESIS_FIGURE_PLACEHOLDER: "
                f"{placeholder_payload['placeholder_id']} | "
                f"{placeholder_payload['scientific_question']} | "
                f"{placeholder_payload['panels'][0]['task']}"
            )
            body_with_placeholder = (
                f"{section_after_placeholder['body'].rstrip()}\n\n{visible_placeholder}"
            )
            status, republished = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/draft",
                method="PUT",
                payload={
                    "section_id": section_after_placeholder["section_id"],
                    "edited_body": body_with_placeholder,
                    "reason": "Human made the figure brief visible and re-approved the current manuscript.",
                    "version_token": section_after_placeholder["version_token"],
                    "actor_type": "simulated_researcher_agent",
                    "actor_label": "generator-figure-human",
                },
            )
            assert status == 200, republished
            republished_section = next(
                row
                for row in republished["sections"]
                if row["section_id"] == "reported-result"
            )
            assert republished_section["body"] == body_with_placeholder
            assert republished_section["decision"]["action"] == "approve"

            figure_markdown_path = (
                "../01_evidence/parses/extracted/prod006-main/images/synthetic_figure.png"
            )
            figure_attribution = (
                f"Source Figure Attribution: {selected_source_figure['figure_id']} | "
                f"{prod006_fixture.SOURCE_ID} | page 1 | Figure 1"
            )
            draft_for_figure_status, draft_for_figure = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/draft",
            )
            assert draft_for_figure_status == 200, draft_for_figure
            section_for_figure = next(
                row
                for row in draft_for_figure["sections"]
                if row["section_id"] == "reported-result"
            )
            body_with_source_figure = (
                f"{section_for_figure['body'].rstrip()}\n\n"
                f"![Synthetic source-bound fixture figure]({figure_markdown_path})\n\n"
                f"{figure_attribution} [evidence:{EVIDENCE_ID}]"
            )
            status, approved_with_figure = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/draft",
                method="PUT",
                payload={
                    "section_id": section_for_figure["section_id"],
                    "edited_body": body_with_source_figure,
                    "reason": "Human selected the synthetic source-bound figure and preserved its attribution.",
                    "version_token": section_for_figure["version_token"],
                    "actor_type": "simulated_researcher_agent",
                    "actor_label": "generator-figure-human-selection",
                },
            )
            assert status == 200, approved_with_figure
            figure_section = next(
                row
                for row in approved_with_figure["sections"]
                if row["section_id"] == "reported-result"
            )
            assert figure_section["body"] == body_with_source_figure
            assert figure_attribution in figure_section["body"]
            assert figure_section["decision"]["action"] == "approve"

            figure_status, figures = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/figures",
            )
            assert figure_status == 200, figures
            assert figures["summary"]["placeholders"] == 1
            assert any(row["state"] == "原论文图" for row in figures["figures"])
            final_status, final = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/final",
            )
            assert final_status == 200, final
            assert final["manuscript_source"] == "authoritative_manuscript"
            export_status, exported = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/export-docx",
                method="POST",
                payload={"release_level": "SELF_REVIEWED_DRAFT"},
            )
            assert export_status == 200, exported
            assert exported["ok"] is True

            release_relative_paths = (
                "05_release/self_reviewed_draft.md",
                "05_release/self_reviewed_draft.docx",
                "05_release/release_snapshot.json",
                "05_release/quality_report.json",
            )
            first_release_bytes = {
                relative: (project / relative).read_bytes()
                for relative in release_relative_paths
            }
            authoritative_bytes = (project / "04_manuscript/manuscript.md").read_bytes()
            assert first_release_bytes["05_release/self_reviewed_draft.md"] == authoritative_bytes
            document = Document(project / "05_release/self_reviewed_draft.docx")
            document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            assert figure_attribution in document_text
            for relative, expected in (
                ("05_release/self_reviewed_draft.md", first_release_bytes["05_release/self_reviewed_draft.md"]),
                ("05_release/self_reviewed_draft.docx", first_release_bytes["05_release/self_reviewed_draft.docx"]),
            ):
                download_status, download_body, download_headers = _raw_request(
                    base_url,
                    f"/file?path={quote(f'{PROJECT_ID}/{relative}', safe='')}" ,
                )
                assert download_status == 200, download_body
                assert download_body == expected
                assert download_headers.get("Content-Type", "")

            current_draft_status, current_draft = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/draft",
            )
            assert current_draft_status == 200, current_draft
            current_section = next(
                row
                for row in current_draft["sections"]
                if row["section_id"] == "reported-result"
            )
            stale_edit = current_section["body"].replace(
                "The limitation stays explicit",
                "The limitation remains explicit",
                1,
            )
            assert stale_edit != current_section["body"]
            edit_status, edited = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/draft",
                method="PUT",
                payload={
                    "section_id": current_section["section_id"],
                    "edited_body": stale_edit,
                    "reason": "Human made a controlled current-version edit before regeneration.",
                    "version_token": current_section["version_token"],
                    "actor_type": "simulated_researcher_agent",
                    "actor_label": "generator-release-currentness",
                },
            )
            assert edit_status == 200, edited

            stale_status, stale_final = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/final",
            )
            assert stale_status == 200, stale_final
            assert stale_final["manuscript_source"] == "authoritative_manuscript"
            assert stale_final["release_status"] == "RELEASE_OUTDATED"
            assert stale_final["release_snapshot"] == {
                "exists": True,
                "matches_authoritative": False,
                "integrity_valid": False,
                "docx_exists": False,
            }
            stale_download_status, stale_download_body, _ = _raw_request(
                base_url,
                f"/file?path={quote(f'{PROJECT_ID}/05_release/self_reviewed_draft.docx', safe='')}",
            )
            assert stale_download_status == 403
            assert b"release DOCX is outdated" in stale_download_body
            assert {
                relative: (project / relative).read_bytes()
                for relative in release_relative_paths
            } == first_release_bytes

            regenerate_status, regenerated = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/export-docx",
                method="POST",
                payload={"release_level": "SELF_REVIEWED_DRAFT"},
            )
            assert regenerate_status == 200, regenerated
            assert regenerated["ok"] is True
            regenerated_bytes = {
                relative: (project / relative).read_bytes()
                for relative in release_relative_paths
            }
            assert regenerated_bytes["05_release/self_reviewed_draft.md"] == (
                project / "04_manuscript/manuscript.md"
            ).read_bytes()
            assert regenerated_bytes["05_release/self_reviewed_draft.md"] != first_release_bytes[
                "05_release/self_reviewed_draft.md"
            ]
            final_after_regenerate_status, final_after_regenerate = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/final",
            )
            assert final_after_regenerate_status == 200
            assert final_after_regenerate["manuscript_source"] == "release_snapshot"
            assert final_after_regenerate["quality_report"]["status"] == "SELF_REVIEWED_DRAFT"
            assert final_after_regenerate["release_snapshot"] == {
                "exists": True,
                "matches_authoritative": True,
                "integrity_valid": True,
                "docx_exists": True,
            }
        finally:
            prod006_fixture._stop_dashboard(server, thread)

        resume_artifact_paths = (
            "04_manuscript/manuscript.md",
            "05_release/self_reviewed_draft.md",
            "05_release/release_snapshot.json",
            "05_release/quality_report.json",
        )
        resume_artifact_refs = [
            {
                "path": relative,
                "sha256": hashlib.sha256((project / relative).read_bytes()).hexdigest(),
            }
            for relative in resume_artifact_paths
        ]
        VersionContext.create(
            {
                "currentness": "current",
                "version_token": "generator-release-current",
                "artifact_refs": resume_artifact_refs,
                "manuscript_source": "release_snapshot",
                "release_level": "SELF_REVIEWED_DRAFT",
            },
            project_root=project,
            project_id=PROJECT_ID,
            version_id="release-v1",
            branch_id="main",
            branch_name="Main",
        )
        current_path = version_context_root(project) / "current.json"
        assert current_path.exists()
        assert (version_context_root(project) / "versions/release-v1.json").is_file()
        assert (version_context_root(project) / "branches/main.json").is_file()
        resume_context = VersionContext.load(project)
        resume_state = resume_context.state()
        resume_node = resume_context.view_version(resume_state.current_version_id)
        resume_payload = {
            "expected_revision": resume_state.revision,
            "node_digest": resume_node.snapshot_digest,
            "version_token": "generator-release-current",
        }

        server, thread, base_url = prod006_fixture._start_dashboard(review_root)
        try:
            resume_status, resumed = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/resume",
                method="POST",
                payload=resume_payload,
            )
            assert resume_status == 200, resumed
            assert resumed["result"] == "RESUMED"
            assert resumed["write_mode"] == "NONE"
            assert resumed["currentness"] == "current"
            assert resumed["version"]["version_id"] == "release-v1"
            assert resumed["version"]["artifact_refs"] == resume_artifact_refs
            assert resumed["revision"] == resume_state.revision

            repeated_status, repeated = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/resume",
                method="POST",
                payload=resume_payload,
            )
            assert repeated_status == 200, repeated
            assert repeated["result"] == "UNCHANGED"
            assert repeated["write_mode"] == "NONE"
            assert repeated["version"] == resumed["version"]
        finally:
            prod006_fixture._stop_dashboard(server, thread)

        # Cold restart must rehydrate the same authoritative current.
        server, thread, base_url = prod006_fixture._start_dashboard(review_root)
        try:
            resumed_after_restart_status, resumed_after_restart = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/resume",
                method="POST",
                payload=resume_payload,
            )
            assert resumed_after_restart_status == 200, resumed_after_restart
            assert resumed_after_restart["result"] == "RESUMED"
            assert resumed_after_restart["write_mode"] == "NONE"
            assert resumed_after_restart["version"] == resumed["version"]
            assert resumed_after_restart["revision"] == resumed["revision"]
            reloaded_status, reloaded_final = _json_request(
                base_url,
                f"/api/project/{PROJECT_ID}/final",
            )
            assert reloaded_status == 200, reloaded_final
            assert reloaded_final["manuscript_source"] == "release_snapshot"
            assert reloaded_final["release_snapshot"] == {
                "exists": True,
                "matches_authoritative": True,
                "integrity_valid": True,
                "docx_exists": True,
            }
            reloaded_download_status, reloaded_download_body, reloaded_download_headers = _raw_request(
                base_url,
                f"/file?path={quote(f'{PROJECT_ID}/05_release/self_reviewed_draft.md', safe='')}",
            )
            assert reloaded_download_status == 200
            assert reloaded_download_body == regenerated_bytes["05_release/self_reviewed_draft.md"]
            assert reloaded_download_headers.get("Content-Type", "")
            reloaded_docx_status, reloaded_docx_body, reloaded_docx_headers = _raw_request(
                base_url,
                f"/file?path={quote(f'{PROJECT_ID}/05_release/self_reviewed_draft.docx', safe='')}",
            )
            assert reloaded_docx_status == 200
            assert reloaded_docx_body == regenerated_bytes["05_release/self_reviewed_draft.docx"]
            assert reloaded_docx_headers.get("Content-Type", "")
            assert {
                relative: (project / relative).read_bytes()
                for relative in release_relative_paths
            } == regenerated_bytes
            assert current_path.exists()
        finally:
            prod006_fixture._stop_dashboard(server, thread)
